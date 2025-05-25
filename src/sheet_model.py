import pandas as pd
from typing import List
from dataclasses import dataclass,field 
from openpyxl.worksheet.worksheet import Worksheet
from datetime import datetime

class Sheet:

    def __init__(self, sheet, header_row_idx=1, data_row_start=None):
        self.sheet : Worksheet = sheet
        self.header_row_idx = header_row_idx
        self.data_row_start= self.header_row_idx+1 if data_row_start is None else data_row_start
        # self.headers = [cell.value for cell in sheet[self.header_row_idx]]
        self.headers =  next(self.sheet.iter_rows(min_row=self.header_row_idx, max_row=self.header_row_idx, values_only=True))
        self.number_of_rows = 0

    def header_idx(self) -> int:
        return self.header_row_idx

    def cell(self, row, column, value):
        self.number_of_rows+=1
        return self.sheet.cell(row=row,column=column, value=value)

    def set_sheet(self, sheet):
        self.sheet=sheet

    def get_self(self):
        return self

    def get_headers(self):
        return self.headers

    def set_data_row_start(self, data_row_start):
        self.data_row_start = data_row_start

    def get_data_row_start(self):
        return self.data_row_start

    def get_sheet(self):
        return self.sheet

    def get_col_idx(self, header):
        return self.headers.index(header)+1

    def get_col_values_for_header(self, header) -> List:
        target_col : int = self.get_col_idx(header)
        col_iter = self.sheet.iter_cols(min_col=target_col,  max_col=target_col, min_row=1, values_only=True)
        return list(next(col_iter))

    def to_dataframe(self, header: str) ->pd.DataFrame:
        col_values =self.get_col_values_for_header(header)
        data : List = col_values[self.data_row_start - 1:]
        return pd.DataFrame({header : data})

    def sheet_name(self):
        return self.sheet.title

    def num_data_rows(self):
        return self.number_of_rows

    def to_full_data_frame(self):
        data= [row for row in self.sheet.iter_rows(min_row=self.data_row_start, values_only=True)]
        return pd.DataFrame(data)


    # def max_row(self):
        # return self.sheet.max_row
