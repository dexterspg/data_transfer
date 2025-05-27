from os import read
import pandas as pd
from openpyxl import load_workbook


def read_sheet_names(file_name : str, header_row :int =1):
    print(f"Reading {file_name}")
    wb = load_workbook(file_name)
    
    for sheet in wb.worksheets:
        print(sheet.title)

        for row in sheet.iter_rows():
            for cell in row:
                print(cell.value)

def read_header_names(file_name : str, header_loc :int =1):
    print(f"Reading {file_name}")
    header_dict={}
    
    xls = pd.ExcelFile(file_name)
    
    for sheetname in xls.sheet_names:
        print(sheetname)
        df = pd.read_excel(file_name, sheetname,  header=header_loc-1, nrows=1)
        header_dict[sheetname]=df.columns.tolist()


    return header_dict

def read_header_names_wb(file_name : str, header_loc :int =1):
    print(f"Reading {file_name}")
    wb = load_workbook(file_name)
    
    header_dict = {}
    for sheetname in wb.sheetnames:
        print(sheetname)
        ws=wb[sheetname]
        header_row = [ cell.value for cell in ws[header_loc]]
        header_dict[sheetname]=header_row

    return header_dict

def put_header_name_into_txt(file_name: str,  header_loc: int = 1, out : str = "header.txt"):
    # header_dict = read_header_names_wb(file_name, header_loc)
    header_dict = read_header_names(file_name, header_loc)
    with open(out, 'w') as f:
        for sheet_name, header_list in header_dict.items():
            f.write(f"Sheet: {sheet_name} :\n")
            for header in header_list:
                f.write(f"{header}\n")

            f.write('\n')
    print(f"saved to {out} file")

def read_header_styles(file_name : str, header_loc :int =1):
    wb=load_workbook(file_name)

    for sheetname in wb.sheetnames:
        print(sheetname)
        ws=wb[sheetname]
        for cell in ws[header_loc]:
            print("Cell Value")
            print(cell.value)
            print("Cell Format")
            print(cell.number_format)
            print("Cell Font")
            print(cell.font)
            print("Cell Fill")
            print(cell.fill)
    

file_name="../poliza_ledger_template.xlsx"
read_header_styles(file_name)






