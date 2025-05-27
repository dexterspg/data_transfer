from openpyxl import load_workbook
from openpyxl.worksheet import worksheet
import pandas as pd
import time


def nombre_formula(header: str, src_df : pd.DataFrame,  *args) -> pd.DataFrame:

    try:
        df=pd.DataFrame(columns=[header])
        df[header] = (
            src_df['Translation Type'] + " " +
            pd.to_datetime(
                src_df['Fiscal Year'].astype(str) + '-' + src_df['Fiscal Period'].astype(str)
            ).dt.strftime('%b-%Y').str.lower()
        )
        return df
    except Exception as e:
        print(e)

def gl_account_split(header: str, src_df: pd.DataFrame, *args) -> pd.DataFrame:
    header_to_index = {
        "compania": 0,
        "Unidad": 1,
        "CC": 2,
        "ubicacion": 3,
        "cuenta": 4,
        "subcuenta": 5,
        "equipo": 6,
        "intercompania": 7,
    }
    
    idx = header_to_index.get(header)
    if idx is None:
        return pd.DataFrame()  # return an empty DataFrame if header not recognized
    
    splitted = src_df["GL Account"].str.split("-", expand=True)
    
    if splitted.shape[1] <= idx:
        col_data = pd.Series([None] * src_df.shape[0])
    else:
        col_data = splitted[idx]
    
    return pd.DataFrame({header: col_data})

formula_mappings = {
    "Nombre": nombre_formula,
    "divisa": lambda header, src_df, *args : pd.DataFrame({ header: src_df["Contract Currency"]}),
    "compania": gl_account_split,
    "Unidad": gl_account_split,
    "CC": gl_account_split,
    "ubicacion": gl_account_split,
    "cuenta": gl_account_split,
    "subcuenta": gl_account_split,
    "equipo": gl_account_split,
    "intercompania": gl_account_split,
    "debito":"",
    "credito":"",
    "debito_convertido":"",
    "credito_convertido":"",
    "descripcion":""
}



def _handle_formula(header: str, src_df,  *args) -> pd.DataFrame:
    formula_func = formula_mappings.get(header)
    
    if callable(formula_func):
        return formula_func(header, src_df, *args)  
    
    return pd.DataFrame()



def mapped_data(source_file: str, template_file: str, updated_file: str, formula_mappings: dict):

    input_header_start=27
    input_data_start=28
    input_df= pd.read_excel(source_file, header=input_header_start-1 , nrows=10)
    template_df = pd.read_excel(template_file) 

    template_wb= load_workbook(template_file, read_only=True)
    template_ws = template_wb.active

    print("Processing source data")
    print(input_df)

    

    for header in template_df.columns:
        if header in formula_mappings.keys():
            result_df :pd.DataFrame = _handle_formula(header, input_df, template_df)
            if not result_df.empty:
                print(result_df)
                template_df[header] = result_df[header]
                print(template_df)

    col_index_map = {
            header: input_df.columns.get_loc(header)
            for header in input_df.columns
        }

    # for row in template_df.itertuples(index=False)
    #     template_ws.cell(
    #         row
    #
    #     )
    #
        
         

    # print(f"Processing worksheet: {template_ws.title} from excel file {template_file}")

def main():

    source_file = "Consolidated Transaction Report.xlsx"
    template_file = "poliza_ledger_template.xlsx"
    updated_file = "poliza_ledger_output.xlsx"

    mapped_data(source_file, template_file, updated_file, formula_mappings)

if __name__ == "__main__":
    main()


