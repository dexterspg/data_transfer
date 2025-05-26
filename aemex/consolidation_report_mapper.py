from openpyxl import load_workbook

def transfer_data(source_file: str, template_file: str, updated_file: str, cell_mappings: dict):

    source_wb = load_workbook(source_file, data_only=True)
    source_ws = source_wb.active
    
    template_wb = load_workbook(template_file)
    template_ws = template_wb.active 
    
    for source_cell, target_cell in cell_mappings.items():
        value = source_ws[source_cell].value
        
        template_ws[target_cell].value = value
        print(f"Copied {value} from {source_cell} to {target_cell}")
    
    template_wb.save(updated_file)
    print(f"Updated report saved as {updated_file}")

cell_mappings = {
    "A1": "B2",
    "B2": "C3",
    "C3": "D4"
}

def main():

    source_file = "source_data.xlsx"
    template_file = "report_template.xlsx"
    updated_file = "updated_report.xlsx"

    transfer_data(source_file, template_file, updated_file, cell_mappings)

if __name__ == "__main__":
    main()


